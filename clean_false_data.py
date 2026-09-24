"""
Script Pembersihan Data Log & Snapshot False-Positive (Data Salah)
===================================================================
Script ini membersihkan:
1. File Excel (.xlsx) & CSV Backup di folder data/:
   - Menghapus baris log dengan plat nomor < 4 karakter (misal: 'G 5', 'T 0', 'W 15')
   - Menghapus baris log dengan status UNREAD
   - Menata ulang nomor urut (No) secara rapi

2. Folder Foto Snapshot (storage/snapshots/ & storage/snapshots_grayscale/):
   - Menghapus file foto snapshot dari plat palsu < 4 karakter / UNREAD
"""

import os
import re
import csv
import openpyxl

DATA_DIR = "data"
SNAPSHOT_DIRS = [
    os.path.join("storage", "snapshots"),
    os.path.join("storage", "snapshots_grayscale")
]


def clean_excel_file(excel_path):
    """Membersihkan baris log palsu dari file Excel sambil menjaga gambar snapshot tetap utuh."""
    if not os.path.exists(excel_path):
        return 0, 0

    try:
        wb = openpyxl.load_workbook(excel_path)
        ws = wb.active

        max_r = ws.max_row
        if max_r <= 1:
            return 0, 0

        # Identifikasi baris yang akan dihapus (dari bawah ke atas agar indeks tidak geser)
        rows_to_delete = []
        for r in range(max_r, 1, -1):
            plate = str(ws.cell(row=r, column=3).value or '').strip().upper()
            plate_clean = re.sub(r"[^A-Z0-9]", "", plate)

            # Validasi plat: minimal 4 karakter (misal: AB1234) dan bukan UNREAD
            if plate_clean == "UNREAD" or len(plate_clean) < 4:
                rows_to_delete.append(r)

        if not rows_to_delete:
            return 0, max_r - 1

        # Hapus baris & sesuaikan gambar anchor di openpyxl
        for r in rows_to_delete:
            target_row_0idx = r - 1
            new_images = []
            for img in getattr(ws, '_images', []):
                img_row = getattr(getattr(img, 'anchor', None), '_from', None)
                if img_row is not None and hasattr(img_row, 'row'):
                    if img_row.row == target_row_0idx:
                        continue
                    elif img_row.row > target_row_0idx:
                        img_row.row -= 1
                        if hasattr(img.anchor, '_to') and hasattr(img.anchor._to, 'row'):
                            img.anchor._to.row -= 1
                new_images.append(img)
            ws._images = new_images

            ws.delete_rows(r)

        # Penataan Ulang Nomor Urut (No) di Kolom A
        for idx, r in enumerate(range(2, ws.max_row + 1), start=1):
            ws.cell(row=r, column=1).value = idx

        wb.save(excel_path)
        return len(rows_to_delete), ws.max_row - 1

    except PermissionError:
        print(f"[WARN] File Excel '{excel_path}' sedang terbuka. Mohon tutup aplikasi Excel!")
        return 0, 0
    except Exception as e:
        print(f"[ERROR] Gagal memproses Excel '{excel_path}': {e}")
        return 0, 0


def clean_csv_file(csv_path):
    """Membersihkan baris log palsu dari file CSV backup."""
    if not os.path.exists(csv_path):
        return 0

    try:
        with open(csv_path, mode="r", encoding="utf-8") as f:
            reader = csv.reader(f)
            rows = list(reader)

        if not rows:
            return 0

        header = rows[0]
        cleaned_rows = [header]
        removed_count = 0

        for row in rows[1:]:
            if not row or len(row) < 2:
                continue

            plate = row[1].strip().upper() if len(row) > 1 else ""
            plate_clean = re.sub(r"[^A-Z0-9]", "", plate)

            if plate_clean == "UNREAD" or len(plate_clean) < 4:
                removed_count += 1
                continue

            cleaned_rows.append(row)

        if removed_count > 0:
            with open(csv_path, mode="w", newline="", encoding="utf-8") as f:
                writer = csv.writer(f)
                writer.writerows(cleaned_rows)

        return removed_count
    except Exception as e:
        print(f"[ERROR] Gagal memproses CSV '{csv_path}': {e}")
        return 0


def clean_snapshot_files():
    """Menghapus file foto snapshot dari plat palsu < 4 karakter / UNREAD."""
    removed_files = 0
    for s_dir in SNAPSHOT_DIRS:
        if not os.path.exists(s_dir):
            continue

        for filename in os.listdir(s_dir):
            if not filename.endswith(".jpg"):
                continue

            match = re.search(r"_([A-Z0-9]+)\.jpg$", filename, re.IGNORECASE)
            if match:
                plate_text = match.group(1).upper()
                if plate_text == "UNREAD" or len(plate_text) < 4:
                    fp = os.path.join(s_dir, filename)
                    try:
                        os.remove(fp)
                        removed_files += 1
                    except Exception:
                        pass
    return removed_files


def main():
    print("[INFO] Memulai Pembersihan Data Log & Snapshot False-Positive...")

    # 1. Membersihkan File Excel & CSV di folder data/
    excel_files = [os.path.join(DATA_DIR, f) for f in os.listdir(DATA_DIR) if f.endswith(".xlsx") and not f.startswith("~$")]
    csv_files = [os.path.join(DATA_DIR, f) for f in os.listdir(DATA_DIR) if f.endswith(".csv")]

    total_excel_removed = 0
    for ef in excel_files:
        rem, total_valid = clean_excel_file(ef)
        print(f"  [EXCEL] '{os.path.basename(ef)}': {rem} baris sampah dihapus, {total_valid} baris valid tersimpan.")
        total_excel_removed += rem

    total_csv_removed = 0
    for cf in csv_files:
        rem = clean_csv_file(cf)
        if rem > 0:
            print(f"  [CSV] '{os.path.basename(cf)}': {rem} baris sampah dihapus.")
            total_csv_removed += rem

    # 2. Membersihkan File Snapshot Foto Sampah
    removed_snapshots = clean_snapshot_files()
    print(f"  [SNAPSHOTS] {removed_snapshots} file foto snapshot sampah dihapus.")

    print("\n[OK] Pembersihan Selesai!")
    print(f"   -> Total Baris Log Sampah Dihapus: {total_excel_removed + total_csv_removed}")
    print(f"   -> Total Foto Snapshot Sampah Dihapus: {removed_snapshots}")


if __name__ == "__main__":
    main()
